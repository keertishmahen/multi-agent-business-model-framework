"""RAG step knowledge uploads (SunMoksha "Knowledge Upload & RAG Integration").

Covers extraction, chunking, per-session isolation, retrieval into the plan
context, curation into every MIA specialist, and the HTTP endpoints. No
network: a deterministic fake embedder and a throwaway Chroma directory.
"""

from __future__ import annotations

import asyncio
import io
import sys
from pathlib import Path

import chromadb
import pytest
from chromadb.config import Settings
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.agents.specialists import SPECIALISTS, build_agent_prompt, filter_rag_context  # noqa: E402
from app.main import app  # noqa: E402
from app.routers import rag as rag_router  # noqa: E402
from app.services import chroma_store as store  # noqa: E402
from app.services import document_ingest as ingest  # noqa: E402
from app.services import rag_service  # noqa: E402

A, B = "session_upload_A", "session_upload_B"
MARKER = "Zebra-42 grant: 40,000 GBP from the Yazd council for rooftop solar"


def _fake_embedding(text: str) -> list[float]:
    vec = [0.0] * 8
    for i, ch in enumerate(text.encode("utf-8")):
        vec[i % 8] += ch
    norm = sum(v * v for v in vec) ** 0.5 or 1.0
    return [v / norm for v in vec]


@pytest.fixture(autouse=True)
def isolated_store(tmp_path, monkeypatch):
    client = chromadb.PersistentClient(
        path=str(tmp_path / "chroma"),
        settings=Settings(anonymized_telemetry=False, allow_reset=True),
    )
    monkeypatch.setattr(store, "_client", client)

    async def fake_embed(text):
        return _fake_embedding(text)

    async def fake_batch(texts):
        return [_fake_embedding(t) for t in texts]

    monkeypatch.setattr(rag_service, "generate_embedding", fake_embed)
    monkeypatch.setattr(rag_service, "generate_embeddings_batch", fake_batch)
    yield


def run(coro):
    return asyncio.run(coro)


def _minimal_pdf(text: str) -> bytes:
    """A one-page PDF with a single text run, assembled by hand."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Extraction & chunking ─────────────────────────────────────────


class TestExtraction:
    def test_txt_and_md(self):
        assert ingest.extract_text("a.txt", MARKER.encode()) == MARKER
        assert "Heading" in ingest.extract_text("notes.MD", b"# Heading\nbody")

    def test_docx_paragraphs_and_tables(self):
        data = _docx(["Market study for refill kiosks", MARKER], [["Cost", "4,000 GBP"]])
        text = ingest.extract_text("study.docx", data)
        assert "Market study" in text and MARKER in text
        assert "Cost | 4,000 GBP" in text

    def test_pdf(self):
        text = ingest.extract_text("plan.pdf", _minimal_pdf("Solar tariff 0.12 per kWh"))
        assert "Solar tariff 0.12 per kWh" in text

    def test_unsupported_type(self):
        with pytest.raises(ingest.UnsupportedDocument):
            ingest.extract_text("sheet.xlsx", b"PK...")

    def test_corrupt_pdf_is_unsupported_not_a_crash(self):
        with pytest.raises(ingest.UnsupportedDocument):
            ingest.extract_text("bad.pdf", b"not a pdf at all")

    def test_too_large(self):
        with pytest.raises(ingest.DocumentTooLarge):
            ingest.extract_text("big.txt", b"x" * (ingest.MAX_FILE_BYTES + 1))

    def test_source_name_is_sanitised(self):
        assert ingest.safe_source_name("C:\\fake\\..\\evil<script>.pdf") == "evil_script_.pdf"


class TestSplitter:
    def test_short_text_is_one_chunk(self):
        assert ingest.split_text("hello world") == ["hello world"]

    def test_chunks_respect_size_and_overlap(self):
        sentences = [f"Sentence number {i} talks about refill stations." for i in range(120)]
        chunks = ingest.split_text(" ".join(sentences))
        assert len(chunks) > 3
        assert all(len(c) <= ingest.CHUNK_SIZE for c in chunks)
        # Consecutive chunks share context (overlap).
        assert any(chunks[i][-40:].split()[-1] in chunks[i + 1] for i in range(len(chunks) - 1))

    def test_unbroken_text_is_hard_split(self):
        chunks = ingest.split_text("a" * 2500)
        assert all(len(c) <= ingest.CHUNK_SIZE for c in chunks)
        assert "".join(chunks).count("a") >= 2500

    def test_empty(self):
        assert ingest.split_text("   ") == []


# ── Ingestion, isolation, retrieval ───────────────────────────────


class TestIngestion:
    def test_ingest_list_and_delete(self):
        docs = run(rag_service.ingest_documents(A, [("grant.txt", MARKER), ("Pasted text", "cost notes")]))
        assert {d["source"] for d in docs} == {"grant.txt", "Pasted text"}
        listed = rag_service.list_documents(A)
        assert len(listed) == 2 and all(d["chunks"] == 1 for d in listed)

        doc_id = next(d["docId"] for d in listed if d["source"] == "grant.txt")
        assert rag_service.delete_document(A, doc_id) == 1
        assert [d["source"] for d in rag_service.list_documents(A)] == ["Pasted text"]

    def test_reupload_is_idempotent(self):
        run(rag_service.ingest_documents(A, [("grant.txt", MARKER)]))
        run(rag_service.ingest_documents(A, [("grant.txt", MARKER)]))
        assert store.count_where(store.Collections.USER_DOCS, {"sessionId": A}) == 1

    def test_sessions_are_isolated(self):
        run(rag_service.ingest_documents(A, [("grant.txt", MARKER)]))
        assert rag_service.list_documents(B) == []
        ctx_b = run(rag_service.retrieve_plan_context(session_id=B, prompt="solar grant"))
        assert ctx_b["uploads"] == []
        # B cannot delete A's document either.
        doc_id = rag_service.list_documents(A)[0]["docId"]
        assert rag_service.delete_document(B, doc_id) == 0
        assert len(rag_service.list_documents(A)) == 1

    def test_chunk_limit(self, monkeypatch):
        monkeypatch.setattr(ingest, "MAX_CHUNKS_PER_SESSION", 2)
        with pytest.raises(ingest.DocumentTooLarge):
            run(rag_service.ingest_documents(A, [("long.txt", "word " * 1000)]))

    def test_uploads_reach_plan_context(self):
        run(rag_service.ingest_documents(A, [("grant.txt", MARKER)]))
        ctx = run(rag_service.retrieve_plan_context(session_id=A, prompt="solar grant funding"))
        assert [u["source"] for u in ctx["uploads"]] == ["grant.txt"]
        assert ctx["uploads"][0]["kind"] == "upload"
        assert ctx["coverageReport"]["uploadsRetrieved"] == 1
        assert ctx["coverageReport"]["uploadsIndexed"] == 1
        assert "upload:grant.txt" in ctx["sources"]
        text = rag_service.format_plan_rag_context(ctx)
        assert MARKER in text and "USER-PROVIDED KNOWLEDGE" in text


class TestCuratedUploadsFeedMia:
    CURATED = [
        {"kind": "bmc", "id": "b1", "blockId": "value_propositions", "text": "value propositions: refills"},
        {"kind": "upload", "id": "u1", "source": "grant.txt", "text": MARKER, "relevance": 0.7},
    ]

    def test_curated_upload_in_plan_text(self):
        ctx = run(rag_service.retrieve_context("plan", "plan", A, {"curatedContext": self.CURATED}))
        assert ctx["uploads"][0]["source"] == "grant.txt"
        assert ctx["coverageReport"]["uploadsRetrieved"] == 1
        assert MARKER in ctx["planRagText"]

    def test_every_specialist_receives_the_upload(self):
        ctx = run(rag_service.retrieve_context("plan", "plan", A, {"curatedContext": self.CURATED}))
        for agent_id, spec in SPECIALISTS.items():
            prompt = build_agent_prompt(
                role=agent_id,
                task="t",
                rag_context=filter_rag_context(ctx, spec.collections),
                user_story="",
                step_id="plan",
            )
            assert MARKER in prompt, agent_id

    def test_excluding_the_upload_keeps_it_out(self):
        ctx = run(rag_service.retrieve_context("plan", "plan", A, {"curatedContext": self.CURATED[:1]}))
        assert MARKER not in ctx["planRagText"]

    def test_routing_lists_upload_for_all_specialists(self):
        assert set(SPECIALISTS) <= set(rag_service.agent_routing()["upload"])


# ── HTTP endpoints ────────────────────────────────────────────────


class TestEndpoints:
    def setup_method(self):
        self.client = TestClient(app)

    def test_upload_files_and_text(self):
        res = self.client.post(
            "/api/rag/upload",
            data={"sessionId": A, "text": "Pasted market note"},
            files=[
                ("files", ("grant.txt", MARKER.encode(), "text/plain")),
                ("files", ("study.docx", _docx(["Market study"]), "application/octet-stream")),
            ],
        )
        assert res.status_code == 200, res.text
        body = res.json()
        assert {d["source"] for d in body["documents"]} == {"grant.txt", "study.docx", "Pasted text"}
        assert body["stored"] == 3

        listed = self.client.get("/api/rag/documents", params={"sessionId": A}).json()["documents"]
        assert len(listed) == 3

        retrieved = self.client.post("/api/rag/retrieve", json={"sessionId": A, "prompt": "grant"}).json()
        assert any(u["source"] == "grant.txt" for u in retrieved["uploads"])
        assert "upload" in retrieved["agentRouting"]

        doc_id = next(d["docId"] for d in listed if d["source"] == "grant.txt")
        assert self.client.delete(f"/api/rag/documents/{doc_id}", params={"sessionId": A}).status_code == 200
        assert self.client.delete(f"/api/rag/documents/{doc_id}", params={"sessionId": A}).status_code == 404

    def test_upload_rejections(self):
        assert self.client.post("/api/rag/upload", data={"sessionId": A}).status_code == 400
        assert self.client.post("/api/rag/upload", data={"text": "x"}).status_code == 400
        bad = self.client.post(
            "/api/rag/upload",
            data={"sessionId": A},
            files=[("files", ("data.xlsx", b"PK", "application/octet-stream"))],
        )
        assert bad.status_code == 415
        too_many = self.client.post(
            "/api/rag/upload",
            data={"sessionId": A},
            files=[("files", (f"f{i}.txt", b"x", "text/plain")) for i in range(6)],
        )
        assert too_many.status_code == 413

    def test_suggest_returns_three_categories(self, monkeypatch):
        async def fake_generate(prompt, **kwargs):
            assert "Business Model Canvas" in prompt and kwargs.get("json_mode")
            return (
                '{"suggestions": ['
                '{"category": "Local market data", "why": "size demand", "examples": ["census"]},'
                '{"category": "Supplier quotes", "why": "costs", "examples": []},'
                '{"category": "Policy incentives", "why": "grants", "examples": ["feed-in tariff"]},'
                '{"category": "extra", "why": "dropped"}]}'
            )

        monkeypatch.setattr(rag_router, "generate_text", fake_generate)
        res = self.client.post("/api/rag/suggest", json={"sessionId": A, "bmc": {"canvas": {}}})
        assert res.status_code == 200
        cats = [s["category"] for s in res.json()["suggestions"]]
        assert cats == ["Local market data", "Supplier quotes", "Policy incentives"]

    def test_suggest_requires_bmc(self):
        assert self.client.post("/api/rag/suggest", json={"sessionId": A}).status_code == 400


# ── Embedding outages are reported, not silent ───────────────────


class TestEmbeddingFailuresAreVisible:
    def test_retrieval_reports_semantic_search_outage(self, monkeypatch):
        async def boom(text):
            raise RuntimeError("Gemini API error (HTTP 402): credits depleted")

        monkeypatch.setattr(rag_service, "generate_embedding", boom)
        ctx = run(rag_service.retrieve_plan_context(session_id=A, prompt="plan"))
        assert "Semantic search unavailable" in ctx["coverageReport"]["retrievalError"]
        assert ctx["knowledge"] == [] and ctx["uploads"] == []

    def test_store_returns_502_when_embedding_fails(self, monkeypatch):
        async def boom(text):
            raise RuntimeError("HTTP 402")

        monkeypatch.setattr(rag_service, "generate_embedding", boom)
        res = TestClient(app).post(
            "/api/context/store",
            json={"sessionId": A, "stepId": "story", "content": "hello", "chunked": True},
        )
        assert res.status_code == 502


# ── Review fixes: hardening of the upload path ───────────────────


class TestUploadHardening:
    def test_long_filename_keeps_its_extension(self):
        long_name = "Market study for the Yazd rooftop solar refill kiosk programme - council " \
            "consultation response and supplier quotes 2026 v3.pdf"
        name = ingest.safe_source_name(long_name)
        assert len(name) <= 120 and name.endswith(".pdf")
        assert "Solar" in ingest.extract_text(name, _minimal_pdf("Solar"))

    def test_docx_decompression_bomb_is_rejected(self):
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/document.xml", "A" * (60 * 1024 * 1024))
        with pytest.raises(ingest.DocumentTooLarge):
            ingest.extract_text("bomb.docx", buf.getvalue())

    def test_not_a_zip_docx_is_unsupported(self):
        with pytest.raises(ingest.UnsupportedDocument):
            ingest.extract_text("fake.docx", b"plain bytes")

    def test_extracted_text_is_capped(self, monkeypatch):
        monkeypatch.setattr(ingest, "MAX_EXTRACTED_CHARS", 100)
        with pytest.raises(ingest.DocumentTooLarge):
            ingest.extract_text("long.txt", b"x" * 101)

    def test_pdf_page_cap(self, monkeypatch):
        monkeypatch.setattr(ingest, "MAX_PDF_PAGES", 0)
        with pytest.raises(ingest.DocumentTooLarge):
            ingest.extract_text("plan.pdf", _minimal_pdf("x"))

    def test_oversized_request_rejected_before_parsing(self):
        res = TestClient(app).post(
            "/api/rag/upload",
            content=b"x",
            headers={
                "content-type": "multipart/form-data; boundary=zzz",
                "content-length": str(ingest.MAX_REQUEST_BYTES + 1),
            },
        )
        assert res.status_code == 413

    def test_bomb_upload_returns_413_over_http(self):
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/document.xml", "A" * (60 * 1024 * 1024))
        res = TestClient(app).post(
            "/api/rag/upload",
            data={"sessionId": A},
            files=[("files", ("bomb.docx", buf.getvalue(), "application/octet-stream"))],
        )
        assert res.status_code == 413
